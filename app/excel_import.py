import math
import re
import unicodedata
import zipfile
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

REQUIRED_COLUMNS = (
    "documento",
    "nombre",
    "programa",
    "semestre",
    "promedio",
    "materias_perdidas",
    "inasistencias",
    "evaluaciones_no_presentadas",
)
_HEADER_ALIASES = {
    "cedula": "documento",
    "identificacion": "documento",
    "numero_de_documento": "documento",
    "nombres": "nombre",
    "materias_perdida": "materias_perdidas",
    "inasistencia": "inasistencias",
    "evaluaciones_no_presentada": "evaluaciones_no_presentadas",
}

MAX_ROWS = 2000
MAX_BYTES = 2 * 1024 * 1024
_DOCUMENTO = re.compile(r"[A-Za-z0-9-]{4,32}")
_FORBIDDEN = re.compile(r"[\x00-\x1f<>]")


def parse_xlsx(data: bytes) -> tuple[list[dict], list[str]]:
    container_error = _validate_container(data)
    if container_error:
        return [], [container_error]
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return [], ["No se pudo leer el archivo. Revisa que sea un .xlsx sin contraseña."]
    try:
        return _read_sheet(workbook)
    finally:
        workbook.close()


def _validate_container(data: bytes) -> str | None:
    if len(data) > MAX_BYTES:
        return "El archivo supera el tamaño máximo de 2 MB."
    if len(data) < 4 or not data.startswith(b"PK"):
        return "El archivo no es un libro de Excel válido."
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            names = [name.lower() for name in archive.namelist()]
            content_types = ""
            if "[content_types].xml" in names:
                original = archive.namelist()[names.index("[content_types].xml")]
                content_types = archive.read(original).decode("utf-8", "ignore")
    except zipfile.BadZipFile:
        return "El archivo no es un libro de Excel válido."
    if "xl/workbook.xml" not in names:
        return "El archivo no es un libro de Excel válido."
    if any("vbaproject" in name for name in names) or "macroenabled" in content_types.lower():
        return "No se aceptan archivos con macros."
    return None


def _read_sheet(workbook) -> tuple[list[dict], list[str]]:
    if not workbook.worksheets:
        return [], ["El archivo no tiene hojas."]
    sheet = workbook.worksheets[0]
    row_iter = sheet.iter_rows(values_only=True)
    headers, header_row_number = _find_headers(row_iter)
    if headers is None:
        return [], ["El archivo no tiene la fila de columnas."]

    errors = []
    for column in REQUIRED_COLUMNS:
        if column == "documento":
            continue
        if column not in headers:
            errors.append(f"Falta la columna obligatoria: {column}.")
        elif headers.count(column) > 1:
            errors.append(f"La columna {column} está repetida.")
    if "documento" in headers and headers.count("documento") > 1:
        errors.append("La columna documento está repetida.")
    if errors:
        return [], errors

    generar_documento = "documento" not in headers
    index = {name: headers.index(name) for name in REQUIRED_COLUMNS if name in headers}
    parsed: list[dict] = []
    seen: dict[str, int] = {}
    data_rows = 0

    for row_number, row in enumerate(row_iter, start=header_row_number + 1):
        values = {name: _cell(row, index[name]) for name in index}
        if generar_documento:
            values["documento"] = _documento_desde_nombre(_text(values.get("nombre")))
        if _row_is_empty(values):
            continue
        data_rows += 1
        if data_rows > MAX_ROWS:
            return [], [f"El archivo supera el máximo de {MAX_ROWS} estudiantes."]
        row_errors, documento, item = _validate_row(row_number, values)
        errors.extend(row_errors)
        if documento:
            previous = seen.get(documento)
            if previous is not None:
                errors.append(
                    f"El documento {documento} está repetido en las filas {previous} y {row_number}."
                )
            else:
                seen[documento] = row_number
        if item is not None and not row_errors:
            parsed.append(item)

    if data_rows == 0 and not errors:
        return [], ["El archivo no tiene estudiantes para cargar."]
    if errors:
        return [], errors
    return parsed, []


def _validate_row(row_number: int, values: dict) -> tuple[list[str], str | None, dict | None]:
    errors: list[str] = []
    documento = _parse_documento(values["documento"], row_number, errors)
    nombre = _parse_text(values["nombre"], "nombre", row_number, errors)
    programa = _parse_text(values["programa"], "programa", row_number, errors)
    semestre = _parse_int(values["semestre"], "semestre", row_number, errors, minimum=1, maximum=20)
    promedio = _parse_promedio(values["promedio"], row_number, errors)
    materias = _parse_int(
        values["materias_perdidas"], "materias_perdidas", row_number, errors, minimum=0, maximum=40
    )
    inasistencias = _parse_int(
        values["inasistencias"], "inasistencias", row_number, errors, minimum=0, maximum=200
    )
    evaluaciones = _parse_int(
        values["evaluaciones_no_presentadas"],
        "evaluaciones_no_presentadas",
        row_number,
        errors,
        minimum=0,
        maximum=40,
    )
    if errors:
        return errors, documento, None
    return [], documento, {
        "documento": documento,
        "nombre": nombre,
        "programa": programa,
        "semestre": semestre,
        "promedio": promedio,
        "materias_perdidas": materias,
        "inasistencias": inasistencias,
        "evaluaciones_no_presentadas": evaluaciones,
    }


def _parse_documento(value, row_number: int, errors: list[str]) -> str | None:
    text = _text(value)
    if text == "":
        errors.append(f"Fila {row_number}: la celda documento está vacía.")
        return None
    if not _DOCUMENTO.fullmatch(text):
        errors.append(
            f"Fila {row_number}: el documento debe tener entre 4 y 32 letras, números o guiones."
        )
        return None
    return text


def _parse_text(value, column: str, row_number: int, errors: list[str]) -> str | None:
    text = _text(value)
    if text == "":
        errors.append(f"Fila {row_number}: la celda {column} está vacía.")
        return None
    if _FORBIDDEN.search(text):
        errors.append(f"Fila {row_number}: {column} contiene caracteres no permitidos.")
        return None
    if not 2 <= len(text) <= 120:
        errors.append(f"Fila {row_number}: {column} debe tener entre 2 y 120 caracteres.")
        return None
    return text


def _parse_promedio(value, row_number: int, errors: list[str]) -> Decimal | None:
    number, kind = _parse_number(value)
    if kind == "vacia":
        errors.append(f"Fila {row_number}: la celda promedio está vacía.")
        return None
    if kind == "texto":
        errors.append(f"Fila {row_number}: en promedio se espera un número.")
        return None
    rounded = round(number, 2)
    if rounded < 0 or rounded > 5:
        errors.append(f"Fila {row_number}: el promedio debe estar entre 0 y 5.")
        return None
    return Decimal(f"{rounded:.2f}")


def _parse_int(value, column: str, row_number: int, errors: list[str], minimum: int, maximum: int) -> int | None:
    number, kind = _parse_number(value)
    if kind == "vacia":
        errors.append(f"Fila {row_number}: la celda {column} está vacía.")
        return None
    if kind == "texto" or abs(number - round(number)) > 1e-9:
        errors.append(f"Fila {row_number}: en {column} se espera un número entero.")
        return None
    parsed = int(round(number))
    if parsed < minimum or parsed > maximum:
        if column == "semestre":
            errors.append(f"Fila {row_number}: el semestre debe ser un entero entre 1 y 20.")
        elif parsed < minimum:
            errors.append(f"Fila {row_number}: en {column} se espera un número entero igual o mayor que cero.")
        else:
            errors.append(f"Fila {row_number}: en {column} el valor es demasiado alto.")
        return None
    return parsed


def _parse_number(value) -> tuple[float | None, str | None]:
    if value is None:
        return None, "vacia"
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == "":
            return None, "vacia"
        cleaned = stripped.replace(" ", "").replace(",", ".")
        if cleaned.count(".") > 1:
            return None, "texto"
        try:
            return float(cleaned), None
        except ValueError:
            return None, "texto"
    if isinstance(value, bool) or isinstance(value, (datetime, date)):
        return None, "texto"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None, "texto"
        return float(value), None
    return None, "texto"


def _find_headers(row_iter) -> tuple[list[str] | None, int]:
    for row_number, row in enumerate(row_iter, start=1):
        if row_number > 15:
            break
        headers = [_canonical_header(cell) for cell in row]
        recognized = [column for column in REQUIRED_COLUMNS if column in headers]
        if len(recognized) >= 4:
            return headers, row_number
    return None, 0


def _canonical_header(value) -> str:
    normalized = _normalize_header(value)
    return _HEADER_ALIASES.get(normalized, normalized)


def _documento_desde_nombre(nombre: str) -> str | None:
    if nombre == "":
        return None
    base = unicodedata.normalize("NFD", nombre.upper())
    base = "".join(char for char in base if unicodedata.category(char) != "Mn" and char.isalnum())
    code = ("AUTO" + base)[:32]
    if len(code) < 4:
        return None
    return code


def _normalize_header(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = re.sub(r"[\s_]+", "_", text).strip("_")
    return text


def _cell(row, position: int):
    if position >= len(row):
        return None
    return row[position]


def _row_is_empty(values: dict) -> bool:
    return all(_text(value) == "" for value in values.values())


def _text(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float):
        if math.isnan(value) or math.isinf(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return str(value).strip()
    if isinstance(value, int):
        return str(value)
    return str(value).strip()
